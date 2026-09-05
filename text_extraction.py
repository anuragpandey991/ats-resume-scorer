"""
text_extraction.py
Extracts and cleans raw text from uploaded resume files (PDF / DOCX).
"""

import re
import io
import pdfplumber
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract text from a PDF file given as bytes.

    x_tolerance=1 (default in pdfplumber is 3) matters a lot here: PDFs with
    justified/tightly-kerned text can have real word gaps smaller than the
    default tolerance, causing pdfplumber to merge separate words into one
    ("Random Forest model using" -> "RandomForestmodelusing"). Lowering the
    tolerance makes it more sensitive to small gaps and fixes this — verified
    directly against a real resume where this was silently corrupting text
    and inflating downstream grammar-check false positives.
    """
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text(x_tolerance=1)
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file given as bytes."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def clean_text(raw_text: str) -> str:
    """Basic cleanup: normalize whitespace, remove weird bullet/icon-font artifacts."""
    if not raw_text:
        return ""
    text = raw_text.replace("\u2022", " ").replace("\uf0b7", " ")  # bullet glyphs
    # Icon-font extraction artifacts: PDFs using icon fonts for contact-info
    # symbols (phone/email/LinkedIn icons) can't be mapped to real characters,
    # so pdfplumber emits placeholders like "(cid:239)". Not real words — strip.
    text = re.sub(r"\(cid:\d+\)", " ", text)
    # Private Use Area unicode chars (U+E000-U+F8FF) are also used by icon
    # fonts and show up as mojibake/garbled single characters.
    text = re.sub(r"[\uE000-\uF8FF]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)          # collapse repeated spaces/tabs
    text = re.sub(r"\n\s*\n+", "\n", text)        # collapse multiple blank lines
    text = text.strip()
    return text


def extract_resume_text(uploaded_file) -> str:
    """
    Given a Streamlit UploadedFile object, detect type by extension
    and return cleaned text. Raises ValueError for unsupported types.
    """
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if filename.endswith(".pdf"):
        raw = extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        raw = extract_text_from_docx(file_bytes)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX resume.")

    cleaned = clean_text(raw)
    if not cleaned:
        raise ValueError(
            "No extractable text found. The file might be a scanned image "
            "(no OCR support in this version) or empty."
        )
    return cleaned


if __name__ == "__main__":
    sample = "Experience\n\u2022 Built REST APIs\n\n\n\u2022 Led   a team of 5"
    print(repr(clean_text(sample)))
